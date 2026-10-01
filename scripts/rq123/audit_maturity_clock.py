#!/usr/bin/env python3
"""Read-only audit of historical Black-76 clocks and merged raw-IV targets.

Print JSON to stdout; never rebuild datasets, modify inputs, or train models.
The alternative-clock price errors are diagnostics, not replacement targets.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.special import ndtr

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))

from wgan_option.merge_support import build_surface_from_params  # noqa: E402

DEFAULT_DATASET = ROOT / (
    "data/processed/rq3/"
    "news_first_vol_surfaces_q097_103_ttm01_38_exact_ttm_v1"
)


def black_prices(futures, strikes, tau, sigma, discount, is_call):
    """Vectorized Black-76, with the discount kept separate from variance time."""
    futures, strikes, tau, sigma, discount = np.broadcast_arrays(
        *[np.asarray(v, dtype=float) for v in (futures, strikes, tau, sigma, discount)]
    )
    if any(not np.isfinite(v).all() for v in (futures, strikes, tau, sigma, discount)):
        raise ValueError("Black inputs must be finite")
    if any((v <= 0).any() for v in (futures, strikes, tau, sigma, discount)):
        raise ValueError("Audited Black inputs must be positive")
    scale = sigma * np.sqrt(tau)
    d1 = np.log(futures / strikes) / scale + 0.5 * scale
    d2 = d1 - scale
    call = discount * (futures * ndtr(d1) - strikes * ndtr(d2))
    put = discount * (strikes * ndtr(-d2) - futures * ndtr(-d1))
    return np.where(is_call, call, put)


def summary(values):
    values = np.asarray(values, dtype=float)
    if values.size == 0 or not np.isfinite(values).all():
        raise ValueError("Cannot summarize empty or non-finite audit values")
    return {
        "min": float(values.min()),
        "mean": float(values.mean()),
        "median": float(np.median(values)),
        "p95": float(np.quantile(values, 0.95)),
        "max": float(values.max()),
    }


def select_accepted(frame):
    """Honor the generation filter; rejected audit rows are not IV targets."""
    selected = frame.loc[
        frame["passes_precalib_filter"].astype(str).str.lower().isin(
            {"true", "1", "1.0"}
        )
    ].copy()
    if selected.empty:
        raise ValueError("No accepted raw observations to audit")
    return selected


def audit_raw_points(frame):
    accepted = select_accepted(frame)
    trade = pd.to_datetime(accepted["trade_datetime_utc"], utc=True)
    expiry = pd.to_datetime(accepted["expiration_datetime_utc"], utc=True)
    tau = (expiry - trade).dt.total_seconds().to_numpy() / (365.0 * 86400.0)
    numeric = {
        name: pd.to_numeric(accepted[name], errors="raise").to_numpy(dtype=float)
        for name in (
            "spot", "strike", "business_days", "implied_vol", "discount_factor",
            "continuous_rate", "price",
        )
    }
    if not all(np.isfinite(value).all() for value in numeric.values()):
        raise ValueError("Accepted observations contain non-finite numbers")
    if (numeric["business_days"] <= 0).any():
        raise ValueError("Accepted observations require positive business days")
    option_type = accepted["option_type"].astype(str).str.upper()
    if not option_type.isin({"CALL", "PUT"}).all():
        raise ValueError("Unexpected option type in accepted observations")
    if not accepted["pricing_model"].eq("black76").all():
        raise ValueError("This audit only supports Black-76 source IVs")
    discount_error = np.abs(
        np.exp(-numeric["continuous_rate"] * tau) - numeric["discount_factor"]
    )
    if np.max(discount_error) > 1e-12:
        raise ValueError("Stored discount factors do not reproduce ACT/365 pricing")

    def prices(clock):
        return black_prices(
            numeric["spot"], numeric["strike"], clock, numeric["implied_vol"],
            numeric["discount_factor"], option_type.eq("CALL").to_numpy(),
        )

    actual_error = np.abs(prices(tau) - numeric["price"])
    if np.max(actual_error) > 1e-8:
        raise ValueError("Stored IVs do not reproduce source prices on ACT/365")
    alternatives = {}
    for denominator in (365, 250):
        proxy_tau = numeric["business_days"] / denominator
        alternatives[f"business_days_div_{denominator}"] = {
            "tau_ratio_to_act365": summary(proxy_tau / tau),
            "tau_difference_above_tolerance": int(np.count_nonzero(np.abs(proxy_tau - tau) > 1e-12)),
            "tau_difference_tolerance": 1e-12,
            "absolute_price_error_fixed_iv_and_discount": summary(
                np.abs(prices(proxy_tau) - numeric["price"])
            ),
        }
    return {
        "rows_total": len(frame),
        "accepted_rows": len(accepted),
        "window_sides": accepted["window_side"].value_counts().to_dict(),
        "act365_repricing_absolute_error": summary(actual_error),
        "act365_discount_absolute_error": summary(discount_error),
        "alternative_clock_diagnostics": alternatives,
    }, accepted


def audit_raw_aggregation(accepted, source):
    """Replay documented aggregation separately from frozen-target checks."""
    # Count each generated snapshot once. A forward JSON entry can reuse the
    # same independently generated backward snapshot as another entry.
    json_snapshots = set()
    selected_counts = {}
    missing_selected_counts = 0
    for directions in source.values():
        for entry in directions.values():
            snapshot = entry["snapshot_time_utc"]
            json_snapshots.add(snapshot)
            count = entry.get("surface_audit", {}).get("selected_option_observations")
            if count is None:
                missing_selected_counts += 1
                continue
            if not isinstance(count, int) or isinstance(count, bool) or count < 0:
                raise ValueError("Invalid JSON selected-option observation count")
            if snapshot in selected_counts and selected_counts[snapshot] != count:
                raise ValueError(f"Conflicting selected-option counts for {snapshot}")
            selected_counts[snapshot] = count
    if missing_selected_counts and selected_counts:
        raise ValueError("Incomplete JSON selected-option observation counts")
    exported = accepted.loc[accepted["window_side"].eq("backward")]
    exported_counts = exported.groupby("target_datetime_utc").size().to_dict()
    if selected_counts:
        deficit_snapshots = {
            snapshot for snapshot, count in selected_counts.items()
            if count > exported_counts.get(snapshot, 0)
        }
        surplus_snapshots = {
            snapshot for snapshot, count in selected_counts.items()
            if count < exported_counts.get(snapshot, 0)
        }
        count_lineage = {
            "json_selected_option_observations_total": sum(selected_counts.values()),
            "count_deficit_snapshot_count": len(deficit_snapshots),
            "count_deficit_rows": sum(
                selected_counts[snapshot] - exported_counts.get(snapshot, 0)
                for snapshot in deficit_snapshots
            ),
            "count_surplus_snapshot_count": len(surplus_snapshots),
            "count_surplus_rows": sum(
                exported_counts[snapshot] - selected_counts[snapshot]
                for snapshot in surplus_snapshots
            ),
        }
    else:
        deficit_snapshots = set()
        count_lineage = {
            "json_selected_option_observations_total": None,
            "count_deficit_snapshot_count": None,
            "count_deficit_rows": None,
            "count_surplus_snapshot_count": None,
            "count_surplus_rows": None,
        }
    numeric = ["business_days", "strike", "percent_strike", "implied_vol", "weight"]
    values = accepted[numeric].astype(float)
    if not np.isfinite(values.to_numpy()).all() or (values <= 0).any().any():
        raise ValueError("Accepted aggregation inputs must be finite and positive")
    if accepted[["target_datetime_utc", "window_side"]].isna().any().any():
        raise ValueError("Accepted aggregation rows lack snapshot lineage")
    slices = {}
    keys = ["target_datetime_utc", "window_side", "business_days", "strike"]
    for key, group in accepted.groupby(keys, sort=False):
        timestamp, side, day, _ = key
        ordered = group.sort_values("implied_vol", kind="stable")
        weights = ordered["weight"].to_numpy(dtype=float)
        index = np.searchsorted(np.cumsum(weights), weights.sum() / 2, side="left")
        iv = float(ordered["implied_vol"].iloc[index])
        strike = float(np.average(group["percent_strike"], weights=group["weight"]))
        slices.setdefault((str(timestamp), str(side), int(day)), []).append((strike, iv))
    for points in slices.values():
        points.sort(key=lambda point: point[0])
    point_count = compared = mismatched = missing = count_errors = 0
    iv_errors, strike_errors, examples = [], [], []
    mismatched_nodes, mismatched_snapshots = set(), set()
    mismatched_on_deficit_snapshots = 0
    for target, directions in source.items():
        for direction, entry in directions.items():
            params = entry.get("surface_params")
            if not params:
                continue
            snapshot = entry["snapshot_time_utc"]
            for day, strikes, vols in zip(
                params["business_days"], params["percent_strikes"], params["implied_vols"]
            ):
                point_count += len(vols)
                # JSON forward uses the later independently generated backward snapshot.
                points = slices.get((snapshot, "backward", int(day)))
                context = {"json_target": target, "direction": direction,
                           "snapshot": snapshot, "business_days": day}
                if points is None:
                    missing += 1
                    if len(examples) < 3:
                        examples.append({**context, "issue": "missing_csv_group"})
                    continue
                if len(points) != len(vols) or len(strikes) != len(vols):
                    count_errors += 1
                    if len(examples) < 3:
                        examples.append({**context, "issue": "point_count_mismatch",
                                         "csv_count": len(points), "json_count": len(vols)})
                    continue
                observed = np.asarray(sorted(zip(strikes, vols)), dtype=float)
                delta = np.asarray(points) - observed
                compared += len(points)
                strike_errors.extend(np.abs(delta[:, 0]))
                iv_errors.extend(np.abs(delta[:, 1]))
                bad = np.any(np.abs(delta) > 1e-12, axis=1)
                mismatched += int(bad.sum())
                if bad.any():
                    mismatched_snapshots.add(snapshot)
                    if snapshot in deficit_snapshots:
                        mismatched_on_deficit_snapshots += int(bad.sum())
                for index in np.flatnonzero(bad):
                    mismatched_nodes.add((snapshot, int(day), int(index)))
                    if len(examples) < 3:
                        examples.append({**context, "issue": "node_mismatch", "point_index": int(index),
                                         "csv_percent_strike_iv": list(points[index]),
                                         "json_percent_strike_iv": observed[index].tolist()})
    return {
        "json_point_count": point_count, "compared_point_count": compared,
        "mismatched_point_count": mismatched, "absolute_tolerance": 1e-12,
        "json_unique_snapshot_count": len(json_snapshots),
        "exported_accepted_backward_rows_total": len(exported),
        "exported_accepted_backward_snapshot_count": len(exported_counts),
        "mismatched_unique_node_count": len(mismatched_nodes),
        "mismatched_unique_snapshot_count": len(mismatched_snapshots),
        "mismatched_point_count_on_deficit_snapshots": (
            mismatched_on_deficit_snapshots if selected_counts else None
        ),
        **count_lineage,
        "missing_group_count": missing, "point_count_error_groups": count_errors,
        "implied_vol_absolute_error": summary(iv_errors) if iv_errors else None,
        "percent_strike_absolute_error": summary(strike_errors) if strike_errors else None,
        "examples": examples,
        "reproduced": bool(point_count and compared == point_count and not (mismatched or missing or count_errors)),
        "interpretation": (
            "Frozen JSON remains authoritative. Differences flag raw-aggregation lineage drift, "
            "not corrected labels or evidence that the maturity clock caused the differences."
        ),
    }


def reconstruct(params, strikes, days, valuation_date, days_in_year):
    surface = build_surface_from_params(
        surface_model="raw", surface_params=params,
        valuation_date=valuation_date, days_in_year=days_in_year,
    )
    return np.asarray(surface.implied_vol_surface(strikes, days), dtype=float).ravel()


def audit_workbook(workbook_path, source, accepted):
    snapshots = set(zip(
        accepted["target_datetime_utc"].astype(str),
        accepted["window_side"].astype(str),
    ))
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rows = workbook["gan_input_ready"].iter_rows(values_only=True)
    headers = next(rows)
    source_row_count = surface_count = value_count = 0
    max_stored_error = max_denominator_error = 0.0
    param_keys, pair_ids = set(), set()
    cached = {}
    try:
        for values in rows:
            row = dict(zip(headers, values))
            if row.get("sample_id") is None:
                continue
            if row["surface_model"] != "raw":
                raise ValueError("This audit only supports raw-IV workbook targets")
            source_row_count += 1
            if not row.get("pair_id"):
                raise ValueError("Workbook row lacks pair_id")
            pair_ids.add(row["pair_id"])
            strikes = json.loads(row["strike_grid"])
            days = json.loads(row["maturity_days_grid"])
            for side, direction in (("current", "backward"), ("target", "forward")):
                params = json.loads(row[f"{side}_surface_param_json"])
                entry = source[row[f"{side}_json_target_timestamp_utc"]][direction]
                if params != entry["surface_params"]:
                    raise ValueError(f"{side} parameters disagree with source JSON")
                # Source forward entries reuse an independently generated backward
                # snapshot: the top-level news timestamp is not the CSV lookup key.
                snapshot = entry["snapshot_time_utc"]
                if (snapshot, "backward") not in snapshots:
                    raise ValueError(f"No accepted raw source for snapshot {snapshot}")
                if snapshot != row[f"{side}_snapshot_time_utc"]:
                    raise ValueError(f"{side} snapshot disagrees with source JSON")
                param_keys.update(params)
                valuation_date = pd.Timestamp(snapshot).date()
                stored = np.asarray(json.loads(row[f"{side}_surface_flat"]), dtype=float)
                cache_key = json.dumps([snapshot, direction, params, strikes, days], sort_keys=True)
                if cache_key not in cached:
                    cached[cache_key] = (
                        reconstruct(params, strikes, days, valuation_date, 250),
                        reconstruct(params, strikes, days, valuation_date, 365),
                    )
                rebuilt, other = cached[cache_key]
                if stored.shape != rebuilt.shape or not np.isfinite(stored).all():
                    raise ValueError(f"Invalid {side} target shape or values")
                stored_error = float(np.max(np.abs(stored - rebuilt)))
                denominator_error = float(np.max(np.abs(rebuilt - other)))
                if stored_error > 1e-12:
                    raise ValueError(f"Stored {side} target does not reproduce raw surface")
                if denominator_error > 1e-12:
                    raise ValueError("Raw interpolation denominator did not cancel")
                max_stored_error = max(max_stored_error, stored_error)
                max_denominator_error = max(max_denominator_error, denominator_error)
                surface_count += 1
                value_count += stored.size
    finally:
        workbook.close()
    if not source_row_count:
        raise ValueError("Workbook has no paired targets")
    return {
        "source_rows": source_row_count,
        "unique_pairs": len(pair_ids),
        "source_endpoint_rows": surface_count,
        "unique_reconstructions": len(cached),
        "surface_values": value_count,
        "source_json_parameter_matches": surface_count,
        "accepted_raw_snapshot_links": surface_count,
        "surface_parameter_keys": sorted(param_keys),
        "max_stored_vs_bus250_reconstruction_error": max_stored_error,
        "max_bus250_vs_bus365_reconstruction_error": max_denominator_error,
        "node_specific_actual_expiry_times_stored": False,
    }


def fingerprint(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return {"path": str(path), "bytes": path.stat().st_size, "sha256": digest.hexdigest()}


def run_audit(dataset):
    inputs = {
        "raw_points": dataset / "shared_surface_inputs/surface-raw-excel-precalib-points.csv",
        "source_surfaces": dataset / "shared_surface_inputs/surface-raw-excel.json",
        "merged_workbook": dataset / "tolerance_05m/merged_vol.xlsx",
    }
    before = {name: fingerprint(path) for name, path in inputs.items()}
    frame = pd.read_csv(inputs["raw_points"], float_precision="round_trip", low_memory=False)
    raw_result, accepted = audit_raw_points(frame)
    source = json.loads(inputs["source_surfaces"].read_text(encoding="utf-8"))
    aggregation = audit_raw_aggregation(accepted, source)
    targets = audit_workbook(inputs["merged_workbook"], source, accepted)
    after = {name: fingerprint(path) for name, path in inputs.items()}
    if before != after:
        raise ValueError("Audit inputs changed during the read-only audit")
    timing_mismatch = raw_result["alternative_clock_diagnostics"]["business_days_div_365"]["tau_difference_above_tolerance"] > 0
    caveats = [name for name, present in (("clock", timing_mismatch), ("lineage", not aggregation["reproduced"])) if present]
    status = "verified_frozen_targets" + ("_with_" + "_and_".join(caveats) + "_caveats" if caveats else "")
    return {
        "status": status,
        "frozen_target_reproduction_verified": True,
        "timing_mismatch": timing_mismatch,
        "raw_aggregation_lineage_caveat": not aggregation["reproduced"],
        "inputs": before,
        "inputs_unchanged": True,
        "raw_points": raw_result,
        "raw_aggregation": aggregation,
        "merged_targets": targets,
        "interpretation": (
            "Source IVs use elapsed ACT/365. Raw grids interpolate sigma^2*q/250 "
            "then divide by q/250; replacing 250 by 365 cancels and is not IV "
            "reannualization. Historical penalties using q/365 therefore use a "
            "business-day proxy, not the original Black pricing time. Alternative "
            "clock price errors hold discount factors fixed and are sensitivity "
            "diagnostics, not evidence that historical source prices were mispriced. "
            "Frozen target reproduction and CSV-to-JSON aggregation lineage are "
            "audited separately; successful target reconstruction does not certify "
            "unified-clock validity or full raw lineage. Rebuilding an actual-expiry grid requires upstream "
            "trade/expiry lineage; existing IV targets and results are preserved."
        ),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset-dir", type=Path, default=DEFAULT_DATASET)
    args = parser.parse_args()
    print(json.dumps(run_audit(args.dataset_dir.resolve()), indent=2, allow_nan=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
