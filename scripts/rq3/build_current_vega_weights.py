"""Freeze current quote-derived Vega weights without changing raw IV labels."""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Mapping

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

import scripts._path_setup  # noqa: E402,F401
import numpy as np
import openpyxl

from film_wgan.support import parse_raw_surface_params, raw_support_mask
from wgan_option.utils.vega_weights import (
    CELL_WEIGHT_MODE,
    SCHEMA_VERSION,
    build_current_vega_grid,
    cell_weight_fingerprint,
    grid_fingerprint,
    load_vega_weight_grid,
    restore_quote_multiplicities,
    sha256_file,
    utc_anchor,
)

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DATASET = ROOT / "data/processed/rq3/news_first_vol_surfaces_q097_103_ttm01_38_exact_ttm_v1"
DEFAULT_OUTPUT = ROOT / "outputs/experiments/rq12_direct_vega_20261007_v1/weights"
DEFAULT_RAW_MANIFEST = ROOT / "data/processed/raw-market-index/raw_2022_2023_london_itm5_relaxed_v2/source_manifest.csv"


def _write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, indent=2, sort_keys=True, allow_nan=False) + "\n")


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise ValueError(f"Cannot write an empty Vega audit: {path}")
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _workbook_current_inputs(workbook_path: Path) -> tuple[dict[str, dict[str, Any]], list[float], list[float], list[dict[str, Any]]]:
    book = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    by_anchor: dict[str, dict[str, Any]] = {}
    sample_audit: list[dict[str, Any]] = []
    grid_q: list[float] | None = None
    grid_days: list[float] | None = None
    try:
        for sheet_name in ("gan_input_ready", "gan_input_atm_ab"):
            if sheet_name not in book.sheetnames:
                continue
            iterator = book[sheet_name].iter_rows(values_only=True)
            columns = next(iterator)
            for values in iterator:
                row = dict(zip(columns, values))
                if row.get("surface_model") != "raw":
                    raise ValueError("Vega sidecars require the frozen raw IV workbook.")
                anchor = utc_anchor(row["current_snapshot_time_utc"])
                if row.get("effective_origin_utc") and utc_anchor(row["effective_origin_utc"]) != anchor:
                    raise ValueError("Workbook effective origin differs from the current snapshot.")
                q = [float(x) for x in json.loads(row["strike_grid"])]
                days = [float(x) for x in json.loads(row["maturity_days_grid"])]
                if grid_q is None:
                    grid_q, grid_days = q, days
                elif q != grid_q or days != grid_days:
                    raise ValueError("Workbook rows use inconsistent surface grids.")
                params = parse_raw_surface_params(row["current_surface_param_json"])
                if anchor in by_anchor and by_anchor[anchor] != params:
                    raise ValueError("One current snapshot has conflicting canonical raw parameters.")
                by_anchor[anchor] = params
                current_support = raw_support_mask(params, strike_grid=q, maturity_days_grid=days)
                target_support = raw_support_mask(
                    parse_raw_surface_params(row["target_surface_param_json"]),
                    strike_grid=q, maturity_days_grid=days,
                )
                # Target metadata is used only to audit the EXISTING loss mask;
                # it is never passed into current quote Vega construction.
                joint_cells = int((current_support & target_support).sum())
                sample_audit.append({
                    "sheet_name": sheet_name, "sample_id": row["sample_id"],
                    "pair_id": row.get("pair_id", ""),
                    "current_snapshot_time_utc": anchor,
                    "target_snapshot_time_utc": row["target_snapshot_time_utc"],
                    "joint_strict_support_cells": joint_cells,
                    "existing_raw_joint_eligible": joint_cells > 0,
                })
    finally:
        book.close()
    if not by_anchor or grid_q is None or grid_days is None:
        raise ValueError("Workbook has no current raw surfaces to cache.")
    return by_anchor, grid_q, grid_days, sample_audit


def build_current_vega_weights(
    dataset_dir: str | Path,
    output_dir: str | Path,
    *,
    tolerance_minutes: int = 5,
    source_manifest_path: str | Path = DEFAULT_RAW_MANIFEST,
) -> dict[str, Any]:
    dataset = Path(dataset_dir).resolve()
    output = Path(output_dir).resolve()
    raw_manifest = Path(source_manifest_path).resolve()
    if (output / "manifest.json").exists():
        raise ValueError("A frozen Vega manifest already exists at the output path; use a new output directory.")
    workbook = dataset / f"tolerance_{tolerance_minutes:02d}m/merged_vol.xlsx"
    precalib = dataset / "shared_surface_inputs/surface-raw-excel-precalib-points.csv"
    canonical, strike_grid, maturity_grid, sample_audit = _workbook_current_inputs(workbook)
    print(f"Caching {len(canonical)} current anchors from both frozen workbook input sheets", flush=True)
    quote_rows: dict[str, list[dict[str, str]]] = defaultdict(list)
    with precalib.open(newline="") as handle:
        for row in csv.DictReader(handle):
            if str(row["passes_precalib_filter"]).lower() != "true":
                continue
            anchor = utc_anchor(row["calibration_datetime_utc"])
            if anchor not in canonical:
                continue
            if row["window_side"] != "backward":
                raise ValueError("Market-index quote provenance must be a trailing backward snapshot.")
            quote_rows[anchor].append(row)
    if set(canonical) != set(quote_rows):
        raise ValueError("Some current snapshots have no accepted quote metadata.")
    counts, raw_audit = restore_quote_multiplicities(
        [row for rows in quote_rows.values() for row in rows], raw_manifest,
        progress=lambda message: print(message, flush=True),
    )
    anchors = sorted(canonical)
    grids: list[np.ndarray] = []
    supports: list[np.ndarray] = []
    anchor_audit: list[dict[str, Any]] = []
    restoration_audit: list[dict[str, Any]] = []
    fingerprints: list[str] = []
    grid_sha = grid_fingerprint(strike_grid, maturity_grid)
    for anchor in anchors:
        grid, audit, quote_audit = build_current_vega_grid(
            quote_rows[anchor], canonical[anchor], counts, strike_grid, maturity_grid,
        )
        fingerprints.append(cell_weight_fingerprint(anchor, grid, grid_sha))
        audit["cell_weight_fingerprint"] = fingerprints[-1]
        grids.append(grid)
        supports.append(raw_support_mask(canonical[anchor], strike_grid=strike_grid, maturity_days_grid=maturity_grid))
        anchor_audit.append(audit)
        restoration_audit.extend(quote_audit)
    for row in sample_audit:
        if row["existing_raw_joint_eligible"]:
            index = anchors.index(row["current_snapshot_time_utc"])
            if not np.any(grids[index] > 0):
                raise ValueError("An existing raw_joint eligible sample has no positive current Vega support.")
    print("All current canonical q/IV buckets reproduced after raw multiplicity restoration", flush=True)
    output.mkdir(parents=True, exist_ok=True)
    npz_path = output / "current_vega_grids.npz"
    np.savez_compressed(
        npz_path,
        current_snapshot_time_utc=np.asarray(anchors),
        vega_weights=np.stack(grids).astype(np.float32),
        current_strict_support=np.stack(supports).astype(bool),
        strike_grid=np.asarray(strike_grid, dtype=np.float64),
        maturity_grid_days=np.asarray(maturity_grid, dtype=np.float64),
        cell_weight_fingerprints=np.asarray(fingerprints),
    )
    _write_csv(output / "anchor_audit.csv", anchor_audit)
    _write_csv(output / "quote_restoration_audit.csv", restoration_audit)
    _write_csv(output / "raw_source_audit.csv", raw_audit)
    _write_csv(output / "sample_coverage_audit.csv", sample_audit)
    qc = {
        "status": "pass", "current_anchor_count": len(anchors),
        "accepted_current_quote_rows": len(restoration_audit),
        "restored_duplicate_occurrences": sum(row["multiplicity"] - 1 for row in restoration_audit),
        "canonical_point_count": sum(row["canonical_point_count"] for row in anchor_audit),
        "maximum_q_reproduction_error": max(row["maximum_q_reproduction_error"] for row in anchor_audit),
        "maximum_iv_reproduction_error": max(row["maximum_iv_reproduction_error"] for row in anchor_audit),
        "raw_source_files_sha_verified": len(raw_audit),
        "current_only_quote_provenance": True,
        "support_outside_weight": 0,
        "zero_current_support_anchors": sum(row["current_strict_support_cells"] == 0 for row in anchor_audit),
        "coverage_sheets": {name: {
            "rows": sum(row["sheet_name"] == name for row in sample_audit),
            "pairs": len({row["pair_id"] for row in sample_audit if row["sheet_name"] == name}),
            "existing_raw_joint_eligible_pairs": len({row["pair_id"] for row in sample_audit if row["sheet_name"] == name and row["existing_raw_joint_eligible"]}),
        } for name in sorted({row["sheet_name"] for row in sample_audit})},
    }
    _write_json(output / "qc_summary.json", qc)
    npz_sha = sha256_file(npz_path)
    manifest = {
        "schema_version": SCHEMA_VERSION, "mode": CELL_WEIGHT_MODE,
        "cell_weight_mode": CELL_WEIGHT_MODE,
        "created_at_utc": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "source_method": "sha_verified_raw_multiplicity_restored_accepted_current_quote_black76_vega",
        "pricing_clock": "quote_expiry_minus_trade_utc_nanoseconds_ACT365",
        "bucket_aggregation": "volume_weighted_median_with_original_raw_duplicate_multiplicity",
        "strike_coordinate": "canonical_volume_weighted_mean_percent_strike",
        "grid_interpolation": "linear_vega_in_percent_strike_then_linear_vega_in_business_days",
        "support_policy": "current_raw_bracket_intersection_v1_outside_zero",
        "loss_normalization": "apply_existing_raw_joint_mask_then_normalize_vega_sum_per_sample",
        "npz_path": npz_path.name, "npz_sha256": npz_sha,
        "weights_path": npz_path.name, "weights_sha256": npz_sha,
        "grid_fingerprint": grid_sha,
        "strike_grid": strike_grid, "maturity_grid_days": maturity_grid,
        "surface_shape": [len(maturity_grid), len(strike_grid)],
        "current_anchor_count": len(anchors), "tolerance_minutes": tolerance_minutes,
        "input_workbook": str(workbook), "input_workbook_sha256": sha256_file(workbook),
        "precalib_csv": str(precalib), "precalib_csv_sha256": sha256_file(precalib),
        "raw_source_manifest": str(raw_manifest), "raw_source_manifest_sha256": sha256_file(raw_manifest),
        "audit_sha256": {path.name: sha256_file(path) for path in sorted(output.glob("*.csv"))},
        "qc_summary_sha256": sha256_file(output / "qc_summary.json"),
        "qc": qc,
    }
    manifest_path = output / "manifest.json"
    _write_json(manifest_path, manifest)
    manifest_sha = sha256_file(manifest_path)
    # Exercise the public loader once before declaring the cache ready.
    for anchor in (anchors[0], anchors[-1]):
        load_vega_weight_grid(manifest_path, anchor, strike_grid, maturity_grid, manifest_sha)
    print(json.dumps({"manifest_path": str(manifest_path), "manifest_sha256": manifest_sha, "qc": qc}, indent=2), flush=True)
    return {"manifest_path": str(manifest_path), "manifest_sha256": manifest_sha, "qc": qc}


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset-dir", type=Path, default=DEFAULT_DATASET)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--tolerance-minutes", type=int, default=5)
    parser.add_argument("--raw-source-manifest", type=Path, default=DEFAULT_RAW_MANIFEST)
    args = parser.parse_args(argv)
    build_current_vega_weights(args.dataset_dir, args.output_dir, tolerance_minutes=args.tolerance_minutes, source_manifest_path=args.raw_source_manifest)


if __name__ == "__main__":
    main()
