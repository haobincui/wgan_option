from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

from scripts.rq123.audit_maturity_clock import (
    audit_raw_points,
    audit_raw_aggregation,
    audit_workbook,
    black_prices,
    fingerprint,
    reconstruct,
)
from wgan_option.market.black76 import black76_price


class MaturityClockAuditTests(unittest.TestCase):
    def raw_frame(self):
        # Friday to Monday spans three calendar days, but only one business day.
        tau = 3 / 365
        sigma, rate = 0.2, 0.04
        discount = float(np.exp(-rate * tau))
        price = black76_price(
            futures_price=100, strike=100, tau=tau, volatility=sigma,
            discount_factor=discount, option_type="CALL",
        )
        return pd.DataFrame([{
            "target_datetime_utc": "2023-01-06T15:00:00Z",
            "window_side": "backward",
            "trade_datetime_utc": "2023-01-06T15:00:00Z",
            "expiration_datetime_utc": "2023-01-09T15:00:00Z",
            "business_days": 1,
            "spot": 100, "strike": 100, "price": price,
            "implied_vol": sigma, "continuous_rate": rate,
            "discount_factor": discount, "option_type": "CALL",
            "pricing_model": "black76", "passes_precalib_filter": True,
            "percent_strike": 1.0, "weight": 1.0,
        }])

    def test_weekend_exposes_proxy_clock_without_rejecting_correct_iv(self):
        report, _ = audit_raw_points(self.raw_frame())
        self.assertLess(report["act365_repricing_absolute_error"]["max"], 1e-12)
        proxy = report["alternative_clock_diagnostics"]["business_days_div_365"]
        self.assertAlmostEqual(proxy["tau_ratio_to_act365"]["mean"], 1 / 3)
        self.assertGreater(proxy["absolute_price_error_fixed_iv_and_discount"]["mean"], 0.1)
        self.assertEqual(proxy["tau_difference_above_tolerance"], 1)

    def test_clock_mismatch_is_measured_instead_of_assumed(self):
        frame = self.raw_frame()
        frame["business_days"] = 3
        report, _ = audit_raw_points(frame)
        proxy = report["alternative_clock_diagnostics"]["business_days_div_365"]
        self.assertEqual(proxy["tau_difference_above_tolerance"], 0)

    def test_reannualization_preserves_both_call_and_put_prices(self):
        old_tau, new_tau, sigma = 3 / 365, 1 / 250, 0.2
        new_sigma = sigma * np.sqrt(old_tau / new_tau)
        for is_call in (True, False):
            old = black_prices(100, 101, old_tau, sigma, 0.99, is_call)
            new = black_prices(100, 101, new_tau, new_sigma, 0.99, is_call)
            self.assertAlmostEqual(float(old), float(new), places=12)

    def test_vectorized_diagnostic_matches_production_calls_and_puts(self):
        futures = np.array([100.0, 105.0, 98.0])
        strikes = np.array([101.0, 100.0, 102.0])
        call_flags = np.array([True, False, False])
        actual = black_prices(futures, strikes, 0.03, 0.2, 0.99, call_flags)
        expected = [
            black76_price(
                futures_price=forward, strike=strike, tau=0.03,
                volatility=0.2, discount_factor=0.99,
                option_type="CALL" if call else "PUT",
            )
            for forward, strike, call in zip(futures, strikes, call_flags)
        ]
        np.testing.assert_allclose(actual, expected, atol=1e-12, rtol=0)

    def test_rejected_rows_are_not_source_targets(self):
        frame = self.raw_frame()
        bad = frame.copy()
        bad["passes_precalib_filter"] = False
        bad["price"] = 999.0
        bad["discount_factor"] = 0.0
        report, accepted = audit_raw_points(pd.concat([frame, bad], ignore_index=True))
        self.assertEqual(report["rows_total"], 2)
        self.assertEqual(len(accepted), 1)

    def test_corrupted_source_price_and_discount_are_rejected(self):
        for field in ("price", "discount_factor"):
            frame = self.raw_frame()
            frame.loc[0, field] += 0.1
            with self.subTest(field=field), self.assertRaises(ValueError):
                audit_raw_points(frame)

    def test_nonfinite_accepted_values_are_rejected(self):
        frame = self.raw_frame()
        frame.loc[0, "continuous_rate"] = np.nan
        with self.assertRaises(ValueError):
            audit_raw_points(frame)

    def workbook_fixture(self, path, *, corrupt_target=False, corrupt_params=False):
        key = "2023-01-06T15:00:00Z"
        future = "2023-01-06T15:05:00Z"
        params = {
            "business_days": [1, 3],
            "implied_vols": [[0.2, 0.22], [0.24, 0.25]],
            "percent_strikes": [[0.98, 1.02], [0.98, 1.02]],
        }
        strikes, days = [0.98, 1.0, 1.02], [1, 2, 3]
        flat = reconstruct(params, strikes, days, date(2023, 1, 6), 250).tolist()
        source = {key: {
            "backward": {"snapshot_time_utc": key, "surface_params": params},
            "forward": {"snapshot_time_utc": future, "surface_params": params},
        }}
        row = {
            "sample_id": "sample-1", "pair_id": "pair-1", "surface_model": "raw",
            "strike_grid": json.dumps(strikes), "maturity_days_grid": json.dumps(days),
        }
        for side, snapshot in (("current", key), ("target", future)):
            row[f"{side}_surface_param_json"] = json.dumps(params)
            row[f"{side}_json_target_timestamp_utc"] = key
            row[f"{side}_snapshot_time_utc"] = snapshot
            row[f"{side}_surface_flat"] = json.dumps(flat)
        if corrupt_target:
            row["target_surface_flat"] = json.dumps([flat[0] + 0.01, *flat[1:]])
        if corrupt_params:
            row["target_surface_param_json"] = json.dumps({**params, "business_days": [1, 4]})
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "gan_input_ready"
        sheet.append(list(row))
        sheet.append(list(row.values()))
        sheet.append(list({**row, "sample_id": "sample-2"}.values()))
        workbook.save(path)
        workbook.close()
        accepted = pd.DataFrame({
            "target_datetime_utc": [key, future], "window_side": ["backward", "backward"],
        })
        return source, accepted

    def test_target_reconstruction_and_forward_snapshot_lineage(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "merged.xlsx"
            source, accepted = self.workbook_fixture(path)
            before = fingerprint(path)
            report = audit_workbook(path, source, accepted)
            self.assertEqual(fingerprint(path), before)
            self.assertEqual(report["source_rows"], 2)
            self.assertEqual(report["unique_pairs"], 1)
            self.assertEqual(report["source_endpoint_rows"], 4)
            self.assertEqual(report["unique_reconstructions"], 2)
            self.assertEqual(report["surface_values"], 36)
            self.assertEqual(report["accepted_raw_snapshot_links"], 4)
            self.assertLess(report["max_bus250_vs_bus365_reconstruction_error"], 1e-12)
            with self.assertRaisesRegex(ValueError, "No accepted raw source"):
                audit_workbook(path, source, accepted.iloc[:1])

    def test_conflicting_repeat_is_checked_even_when_reconstruction_is_cached(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "merged.xlsx"
            source, accepted = self.workbook_fixture(path)
            workbook = load_workbook(path)
            sheet = workbook.active
            headers = [cell.value for cell in sheet[1]]
            column = headers.index("target_surface_flat") + 1
            values = json.loads(sheet.cell(3, column).value)
            values[0] += 0.01
            sheet.cell(3, column, json.dumps(values))
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "Stored target"):
                audit_workbook(path, source, accepted)

    def test_corrupted_target_and_source_params_are_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "merged.xlsx"
            for corrupt_field in ("corrupt_target", "corrupt_params"):
                source, accepted = self.workbook_fixture(path, **{corrupt_field: True})
                with self.subTest(field=corrupt_field), self.assertRaises(ValueError):
                    audit_workbook(path, source, accepted)

    def test_aggregation_weighted_median_and_forward_snapshot(self):
        raw = self.raw_frame().iloc[0].to_dict()
        origin, future = raw["target_datetime_utc"], "2023-01-06T15:05:00Z"
        points = []
        for snapshot in (origin, future):
            points.extend([
                {**raw, "target_datetime_utc": snapshot, "implied_vol": 0.1,
                 "percent_strike": 0.9, "weight": 1.0},
                {**raw, "target_datetime_utc": snapshot, "implied_vol": 0.3,
                 "percent_strike": 1.1, "weight": 1.0},
            ])
        source = {origin: {
            side: {"snapshot_time_utc": snapshot, "surface_params": {
                "business_days": [1], "percent_strikes": [[1.0]], "implied_vols": [[0.1]],
            }} for side, snapshot in (("backward", origin), ("forward", future))
        }}
        report = audit_raw_aggregation(pd.DataFrame(points), source)
        self.assertTrue(report["reproduced"])
        self.assertEqual(report["json_point_count"], 2)
        self.assertEqual(report["mismatched_point_count"], 0)
        source[origin]["forward"]["surface_params"]["implied_vols"] = [[0.3]]
        report = audit_raw_aggregation(pd.DataFrame(points), source)
        self.assertFalse(report["reproduced"])
        self.assertEqual(report["mismatched_point_count"], 1)
        self.assertAlmostEqual(report["implied_vol_absolute_error"]["max"], 0.2)

    def test_aggregation_missing_groups_are_explicit_caveats(self):
        with tempfile.TemporaryDirectory() as temporary:
            source, _ = self.workbook_fixture(Path(temporary) / "merged.xlsx")
            report = audit_raw_aggregation(self.raw_frame(), source)
            self.assertFalse(report["reproduced"])
            self.assertEqual(report["json_point_count"], 8)
            self.assertEqual(report["point_count_error_groups"], 1)
            self.assertEqual(report["missing_group_count"], 3)

    def test_observation_counts_deduplicate_reused_json_snapshots(self):
        raw = self.raw_frame().iloc[0].to_dict()
        origin, future = raw["target_datetime_utc"], "2023-01-06T15:05:00Z"
        accepted = pd.DataFrame([
            raw,
            {**raw, "target_datetime_utc": future},
            {**raw, "target_datetime_utc": future, "window_side": "forward",
             "implied_vol": 0.3},
        ])

        def entry(snapshot, iv, selected):
            return {
                "snapshot_time_utc": snapshot,
                "surface_audit": {"selected_option_observations": selected},
                "surface_params": {
                    "business_days": [1], "percent_strikes": [[1.0]],
                    "implied_vols": [[iv]],
                },
            }

        source = {
            origin: {
                "backward": entry(origin, 0.2, 2),
                "forward": entry(future, 0.3, 2),
            },
            future: {"backward": entry(future, 0.3, 2)},
        }
        report = audit_raw_aggregation(accepted, source)
        self.assertEqual(report["json_unique_snapshot_count"], 2)
        self.assertEqual(report["json_selected_option_observations_total"], 4)
        self.assertEqual(report["exported_accepted_backward_rows_total"], 2)
        self.assertEqual(report["count_deficit_snapshot_count"], 2)
        self.assertEqual(report["count_deficit_rows"], 2)
        self.assertEqual(report["mismatched_point_count"], 2)
        self.assertEqual(report["mismatched_unique_node_count"], 1)
        self.assertEqual(report["mismatched_unique_snapshot_count"], 1)
        self.assertEqual(report["mismatched_point_count_on_deficit_snapshots"], 2)

        source[future]["backward"]["surface_audit"]["selected_option_observations"] = 3
        with self.assertRaisesRegex(ValueError, "Conflicting selected-option counts"):
            audit_raw_aggregation(accepted, source)


if __name__ == "__main__":
    unittest.main()
